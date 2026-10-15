"""Your full CV, pasted once and searched whenever it's needed.

The profile is a summary: a few lines, lists of skills and employers. A CV
says far more — what you did at each job, the projects, the numbers — and
forms ask for exactly that ("describe a time you…", "what did you build at
X?"). Paste the whole CV here and:

  * it's split into sections and indexed, so drafting a letter or answering a
    form question pulls the parts of the CV that matter to that question
    (retrieval, not the whole CV every time);
  * empty profile fields are filled from it — headline, skills, technologies,
    employers, achievements — so the profile isn't typed twice;
  * the claims check accepts what's in the CV as yours: a draft that names a
    tool from your CV isn't held as invented.

Search is BM25 over words, with nothing to install and nothing sent anywhere.
Stored in the job data folder: jobscout/cv/cv.txt and chunks.json.
"""
from __future__ import annotations

import json
import math
import re
import time
from pathlib import Path

from . import config

MAX_CHARS = 120_000          # a long CV is ~15k; this stops a pasted book
CHUNK = 700                  # characters per section, roughly a paragraph
_WORD = re.compile(r"[a-z0-9][a-z0-9+#.\-]*")
_STOP = set("""a an and are as at be by for from has have i in is it its of on
or our that the their this to was were will with you your my me we us he she
they them his her not no but if so than then there these those which who what
when where how all any can into out up""".split())
# lines that open a section in most CVs
_HEADING = re.compile(
    r"^\s*(profile|summary|about|experience|employment|work history|career|"
    r"education|qualifications?|certifications?|skills|technical skills|"
    r"projects?|achievements|awards|publications|languages|interests|"
    r"references|volunteer|training|courses?|top skills|contact|honors-awards)\b[\s:]*$", re.I)
_LI_URL = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[A-Za-z0-9_%-]+/?", re.I)
_LI_PAGE = re.compile(r"^\s*Page \d+ of \d+\s*$", re.I | re.M)


def looks_like_linkedin(text: str) -> bool:
    """LinkedIn's "Save to PDF" of a profile: its link, its page footers and
    its sidebar headings."""
    t = text or ""
    return bool(_LI_URL.search(t)) and (bool(_LI_PAGE.search(t)) or "Top Skills" in t)


def from_linkedin(text: str) -> dict:
    """Tidy a LinkedIn profile PDF's text, and pick out what it states plainly:
    the profile link, the name and headline at the top of the main column,
    and the location under them. Nothing is guessed."""
    t = _LI_PAGE.sub("", text or "")
    t = re.sub(r"\s*\(LinkedIn\)", "", t)
    url = _LI_URL.search(t)
    out = {"text": t.strip(), "linkedin": "", "full_name": "", "headline": "", "location": ""}
    if url:
        u = url.group(0)
        out["linkedin"] = u if u.startswith("http") else "https://" + u.lstrip("/")
    lines = [l.strip() for l in t.split("\n")]
    # the main column starts with the name, right before "Summary" / "Experience"
    # by a few lines: name, headline, location
    for i, l in enumerate(lines):
        if l.lower() in ("summary", "experience") and i >= 3:
            block = [x for x in lines[max(0, i - 4):i] if x]
            if len(block) >= 3:
                out["full_name"], out["headline"], out["location"] = block[-3], block[-2], block[-1]
            break
    if out["full_name"] and (len(out["full_name"].split()) > 5 or any(ch.isdigit() for ch in out["full_name"])):
        out.update(full_name="", headline="", location="")
    return out



def _dir() -> Path:
    d = config.AGENT_HOME / "jobscout" / "cv"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _words(text: str) -> list:
    return [w.strip(".-") for w in _WORD.findall((text or "").lower())
            if w.strip(".-") and w not in _STOP and len(w) > 1]


def chunk(text: str) -> list:
    """Split a CV into sections of about CHUNK characters, each carrying the
    heading it sits under, so "Standard Bank — 2019–2024" travels with the
    bullet points beneath it."""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    out, heading, buf = [], "", []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            out.append({"heading": heading, "text": body})
        buf.clear()

    for line in text.split("\n"):
        if _HEADING.match(line) or (line.isupper() and 2 < len(line.strip()) < 40):
            flush()
            heading = line.strip().strip(":").title()
            continue
        if not line.strip() and sum(len(x) for x in buf) >= CHUNK * 0.6:
            flush()
            continue
        buf.append(line)
        if sum(len(x) for x in buf) >= CHUNK:
            flush()
    flush()
    for i, c in enumerate(out):
        c["id"] = i
    return out


def extract_text(filename: str, data: bytes) -> dict:
    """The text of an uploaded CV: PDF, Word (.docx), or plain text."""
    import io
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf") or data[:5] == b"%PDF-":
            from pypdf import PdfReader
            pages = [p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages]
            text = "\n\n".join(pages)
            if len(text.strip()) < 200:
                return {"ok": False, "error": "That PDF has almost no text in it — "
                        "it's probably a scan. Export the CV from Word as a PDF, "
                        "or paste its text instead."}
            return {"ok": True, "text": text}
        if name.endswith(".docx") or data[:2] == b"PK":
            import docx
            doc = docx.Document(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs]
            for t in doc.tables:                 # CV templates love tables
                for row in t.rows:
                    lines.append("  ".join(c.text for c in row.cells))
            return {"ok": True, "text": "\n".join(lines)}
        if name.endswith(".doc"):
            return {"ok": False, "error": "Old Word (.doc) files can't be read — "
                    "save it as .docx or PDF and upload that."}
        return {"ok": True, "text": data.decode("utf-8", "replace")}
    except Exception as exc:
        return {"ok": False, "error": f"Couldn't read that file ({type(exc).__name__})."}


def save(text: str, name: str = "") -> dict:
    """Store the pasted CV and index it. Replaces the previous one."""
    text = (text or "").strip()[:MAX_CHARS]
    linkedin = {}
    if looks_like_linkedin(text):
        linkedin = from_linkedin(text)
        text = linkedin["text"]
    if len(text) < 200:
        return {"ok": False, "error": "That's too short to be a CV — paste the "
                                      "whole thing, every section."}
    d = _dir()
    (d / "cv.txt").write_text(text, "utf-8")
    parts = chunk(text)
    (d / "chunks.json").write_text(json.dumps(
        {"name": name or "CV", "at": time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime()),
         "chunks": parts}, indent=1), "utf-8")
    out = {"ok": True, "chars": len(text), "sections": len(parts)}
    if linkedin:
        out["linkedin"] = {k: linkedin[k] for k in ("linkedin", "full_name", "headline", "location")}
    return out


def text() -> str:
    try:
        return (_dir() / "cv.txt").read_text("utf-8")
    except Exception:
        return ""


def status() -> dict:
    try:
        d = json.loads((_dir() / "chunks.json").read_text("utf-8"))
    except Exception:
        return {"present": False}
    return {"present": True, "name": d.get("name", "CV"), "at": d.get("at", ""),
            "sections": len(d.get("chunks") or []), "chars": len(text())}


def clear() -> dict:
    for f in ("cv.txt", "chunks.json"):
        try:
            (_dir() / f).unlink()
        except FileNotFoundError:
            pass
    return {"ok": True}


def search(query: str, k: int = 4) -> list:
    """The k sections that best answer `query` (BM25). [] without a CV."""
    try:
        parts = json.loads((_dir() / "chunks.json").read_text("utf-8"))["chunks"]
    except Exception:
        return []
    q = set(_words(query))
    if not parts or not q:
        return []
    docs = [_words(p["heading"] + " " + p["text"]) for p in parts]
    n = len(docs)
    avg = sum(len(d) for d in docs) / max(1, n)
    df = {w: sum(1 for d in docs if w in d) for w in q}
    scored = []
    for p, d in zip(parts, docs):
        s = 0.0
        for w in q:
            tf = d.count(w)
            if not tf:
                continue
            idf = math.log(1 + (n - df[w] + 0.5) / (df[w] + 0.5))
            s += idf * tf * 2.2 / (tf + 1.2 * (0.25 + 0.75 * len(d) / avg))
        if s > 0:
            scored.append((s, p))
    scored.sort(key=lambda x: -x[0])
    return [{"heading": p["heading"], "text": p["text"], "score": round(s, 2)}
            for s, p in scored[:k]]


def excerpts(query: str, k: int = 4, limit: int = 3500) -> str:
    """Search results as one block to hand an engine."""
    out, used = [], 0
    for h in search(query, k):
        piece = (f"[{h['heading']}]\n" if h["heading"] else "") + h["text"]
        if used + len(piece) > limit:
            break
        out.append(piece)
        used += len(piece)
    return "\n---\n".join(out)


_EXTRACT_SYSTEM = (
    "You read a candidate's CV and return the facts in it as JSON, copying "
    "what the CV says — never adding, never guessing. Omit anything the CV "
    "does not state. Return ONLY JSON with any of these keys: "
    "{\"full_name\": str, \"email\": str, \"phone\": str, \"location\": str, "
    "\"linkedin\": str, \"headline\": str, \"summary\": str (2-3 sentences "
    "in the candidate's own words from the CV), \"skills\": [str], "
    "\"technologies\": [str], \"employers\": [str], \"achievements\": [str] "
    "(concrete, as written), \"certifications\": [str], "
    "\"years_experience\": {area: years} only where the CV states or dates "
    "make it plain}.")


def to_profile(brain, model=None) -> dict:
    """Fill the profile's empty fields from the CV. Fields you've already
    written are never overwritten; lists are added to, not replaced."""
    from . import jobscout
    cv = text()
    if not cv:
        return {"ok": False, "error": "No CV pasted yet."}
    if brain is None:
        return {"ok": False, "error": "No engine available to read the CV."}
    try:
        data = jobscout._engine_json(brain, cv[:24000], _EXTRACT_SYSTEM, model)
    except Exception as exc:
        return {"ok": False, "error": jobscout._explain_engine_error(
            f"{type(exc).__name__}: {exc}")}
    prof = jobscout.profile()
    changes, filled = {}, []
    for key, val in (data or {}).items():
        if val in (None, "", [], {}):
            continue
        have = prof.get(key)
        if isinstance(val, list):
            old = list(have) if isinstance(have, list) else []
            low = {str(x).strip().lower() for x in old}
            new = [str(v).strip() for v in val
                   if str(v).strip() and str(v).strip().lower() not in low]
            if new:
                changes[key] = old + new
                filled.append(f"{key} (+{len(new)})")
        elif isinstance(val, dict):
            old = dict(have) if isinstance(have, dict) else {}
            add = {k: v for k, v in val.items() if k not in old}
            if add:
                changes[key] = {**old, **add}
                filled.append(key)
        elif not str(have or "").strip():
            changes[key] = str(val).strip()
            filled.append(key)
    if changes:
        jobscout.save_profile(changes)
    return {"ok": True, "filled": filled}
